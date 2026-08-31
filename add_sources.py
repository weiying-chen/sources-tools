#!/home/weiying/python/word/.venv/bin/python
"""Add queued source DOCX entries to the current programme's master table."""

from __future__ import annotations

import argparse
import re
from copy import deepcopy
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.table import _Cell
from docx.text.run import Run


DATE_RE = re.compile(r"(20\d{6})(?=\.docx$)")
NUMBER_RE = re.compile(r"^\s*(\d+)")


@dataclass(frozen=True)
class QueuedEntry:
    path: Path
    date: int
    title: str
    url: str
    timestamp: str


def nonempty_paragraphs(path: Path) -> list[str]:
    document = Document(path)
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def read_queued_entries(queued_dir: Path) -> list[QueuedEntry]:
    entries: list[QueuedEntry] = []
    for path in sorted(queued_dir.glob("*.docx")):
        if path.name.startswith("~$") or path.name.endswith(":Zone.Identifier"):
            continue
        date_match = DATE_RE.search(path.name)
        paragraphs = nonempty_paragraphs(path)
        if not date_match:
            raise RuntimeError(f"Queued filename has no YYYYMMDD date: {path.name}")
        if len(paragraphs) < 3:
            raise RuntimeError(f"Queued DOCX needs title, URL, and timestamp: {path.name}")
        title, url, timestamp = paragraphs[:3]
        if not url.startswith(("https://youtu.be/", "https://www.youtube.com/")):
            raise RuntimeError(f"Second paragraph is not a YouTube URL: {path.name}")
        entries.append(
            QueuedEntry(
                path=path,
                date=int(date_match.group(1)),
                title=title,
                url=url,
                timestamp=timestamp,
            )
        )
    return sorted(entries, key=lambda entry: (entry.date, entry.path.name))


def set_paragraph_text(paragraph, text: str) -> None:
    text_nodes = paragraph._p.xpath(".//w:t")
    if text_nodes:
        text_nodes[0].text = text
        for node in text_nodes[1:]:
            node.text = ""
    else:
        paragraph.add_run(text)
    for run_element in paragraph._p.xpath(".//w:r"):
        run = Run(run_element, paragraph)
        run.font.name = "Calibri"
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "新細明體")


def clear_cell(cell) -> None:
    for paragraph in cell.paragraphs:
        set_paragraph_text(paragraph, "")


def numeric_row_value(row) -> int | None:
    if not row.cells:
        return None
    match = NUMBER_RE.match(row.cells[0].text)
    return int(match.group(1)) if match else None


def build_output_path(master: Path, in_place: bool) -> Path:
    if in_place:
        return master
    return master.with_name(f"{master.stem}_updated{master.suffix}")


def discover_master(project_dir: Path) -> Path:
    candidates = sorted(
        path
        for path in project_dir.glob("*.docx")
        if not path.name.startswith("~$") and not path.stem.endswith("_updated")
    )
    if len(candidates) == 1:
        return candidates[0]
    if not candidates:
        raise SystemExit(
            f"No master DOCX found in {project_dir}. Put exactly one master DOCX "
            "in the programme folder or pass --master."
        )
    raise SystemExit(
        "Multiple possible master DOCX files found; pass --master: "
        + ", ".join(path.name for path in candidates)
    )


def remove_blank_data_rows(table) -> int:
    removed = 0
    for row in list(table.rows[1:]):
        if len(row.cells) > 1 and not row.cells[1].text.strip():
            table._tbl.remove(row._tr)
            removed += 1
    return removed


def discover_queued_dir(project_dir: Path) -> Path:
    queued = project_dir / "queued"
    output = project_dir / "output"
    if queued.is_dir() and any(queued.glob("*.docx")):
        return queued
    if output.is_dir() and any(output.glob("*.docx")):
        return output
    return queued


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--master",
        type=Path,
        help="master DOCX (defaults to the only DOCX in the current programme folder)",
    )
    parser.add_argument(
        "--queued-dir",
        type=Path,
        help="source files folder (defaults to ./queued, or ./output when queued is empty)",
    )
    parser.add_argument(
        "--copy",
        dest="in_place",
        action="store_false",
        help="write MASTER_updated.docx instead of updating the master in place",
    )
    parser.add_argument("--in-place", dest="in_place", action="store_true", help=argparse.SUPPRESS)
    parser.set_defaults(in_place=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    project_dir = Path.cwd()
    master = args.master if args.master else discover_master(project_dir)
    queued_dir = args.queued_dir if args.queued_dir else discover_queued_dir(project_dir)
    if not master.is_file():
        raise SystemExit(f"Missing master DOCX: {master}")
    if not queued_dir.is_dir():
        raise SystemExit(f"Missing queued directory: {queued_dir}")

    entries = read_queued_entries(queued_dir)
    if not entries:
        print("No queued DOCX files found.")
        return 0

    document = Document(master)
    if not document.tables:
        raise SystemExit("Master DOCX has no table.")
    table = document.tables[0]
    removed_blanks = remove_blank_data_rows(table)
    if len(table.rows) < 2 or len(table.rows[1].cells) < 2:
        raise SystemExit("Master table has no usable data-row template.")

    existing_titles = {row.cells[1].paragraphs[0].text.strip() for row in table.rows[1:] if row.cells[1].paragraphs}
    pending = [entry for entry in entries if entry.title not in existing_titles]
    skipped = len(entries) - len(pending)
    if not pending and not removed_blanks:
        print(f"No new queued entries found; skipped {skipped} existing row(s).")
        return 0

    sequence_values = [value for row in table.rows for value in [numeric_row_value(row)] if value is not None]
    next_sequence = max(sequence_values, default=0) + 1
    template_row = table.rows[1]
    header_row = table.rows[0]

    planned: list[tuple[int, QueuedEntry]] = []
    for entry in pending:
        sequence = next_sequence
        next_sequence += 1
        planned.append((sequence, entry))
        print(f"{sequence}\t{entry.title}")

        if args.dry_run:
            continue

        new_tr = deepcopy(template_row._tr)
        header_row._tr.addnext(new_tr)
        physical_cells = [_Cell(tc, table) for tc in new_tr.tc_lst]
        set_paragraph_text(physical_cells[0].paragraphs[0], str(sequence))

        target = physical_cells[1]
        while len(target.paragraphs) < 3:
            target.add_paragraph("")
        set_paragraph_text(target.paragraphs[0], entry.title)
        set_paragraph_text(target.paragraphs[1], entry.url)
        set_paragraph_text(target.paragraphs[2], entry.timestamp)
        for paragraph in target.paragraphs[3:]:
            set_paragraph_text(paragraph, "")

        for cell in physical_cells[2:]:
            clear_cell(cell)

    output = build_output_path(master, args.in_place)
    if args.dry_run:
        print(
            f"Dry run: would add {len(planned)} row(s), remove {removed_blanks} "
            f"blank row(s), skip {skipped}, output {output}"
        )
        return 0

    document.save(output)
    # Reopen through python-docx to catch malformed package/XML output.
    validated = Document(output)
    if not validated.tables or len(validated.tables[0].rows) != len(table.rows):
        output.unlink(missing_ok=True)
        raise RuntimeError("Saved DOCX failed table validation.")
    print(
        f"Wrote {output} with {len(planned)} new row(s), removed "
        f"{removed_blanks} blank row(s); skipped {skipped}."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
