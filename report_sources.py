#!/home/weiying/python/word/.venv/bin/python
"""Report pending PM translation/editing counts and copy them to the clipboard."""

from __future__ import annotations

import argparse
import re
import subprocess
import tomllib
from dataclasses import dataclass
from pathlib import Path

from docx import Document


SCRIPT_DIR = Path(__file__).resolve().parent
ROOT = Path.cwd()
CONFIG_PATH = SCRIPT_DIR / "status_people.toml"
PROGRAMMES = (
    ("大愛醫生館", ROOT / "all-about-health/4大愛醫生館(菊芬雲端).docx"),
    ("大愛真健康", ROOT / "easy-fitness/4大愛真健康.docx"),
)
TIMESTAMP_RE = re.compile(r"^\d{1,2}:\d{2}")


@dataclass(frozen=True)
class PersonPolicy:
    name: str
    edit_required: bool


def normalize_name(value: str) -> str:
    return re.sub(r"\s+", "", value).casefold()


def load_config(path: Path) -> tuple[dict[str, PersonPolicy], int]:
    with path.open("rb") as config_file:
        data = tomllib.load(config_file)
    policies: dict[str, PersonPolicy] = {}
    for item in data.get("people", []):
        policy = PersonPolicy(str(item["name"]), bool(item["edit_required"]))
        for alias in item.get("aliases", [policy.name]):
            key = normalize_name(str(alias))
            if key in policies:
                raise SystemExit(f"Duplicate person alias in {path}: {alias}")
            policies[key] = policy
    ready_target = int(data.get("report", {}).get("translation_ready_target", 3))
    if ready_target < 1:
        raise SystemExit("translation_ready_target must be at least 1")
    return policies, ready_target


def has_english_details(cell) -> bool:
    paragraphs = [paragraph.text.strip() for paragraph in cell.paragraphs if paragraph.text.strip()]
    url_index = next((index for index, text in enumerate(paragraphs) if "youtu" in text.lower()), None)
    if url_index is None:
        return False
    for text in paragraphs[url_index + 1 :]:
        if TIMESTAMP_RE.match(text):
            continue
        if re.search(r"[A-Za-z]{3}", text):
            return True
    return False


def active_rows(master: Path):
    document = Document(master)
    if not document.tables:
        raise SystemExit(f"Master DOCX has no table: {master}")
    table = document.tables[0]
    if len(table.columns) < 4:
        raise SystemExit(f"Master table has too few columns: {master}")
    for row in table.rows[1:]:
        if len(row.cells) < 4 or not row.cells[1].text.strip():
            continue
        # The first production status marks the end of the current backlog.
        if any(cell.text.strip() for cell in row.cells[3:]):
            break
        yield row


def count_programme(master: Path, policies: dict[str, PersonPolicy]) -> tuple[int, int]:
    if not master.is_file():
        raise SystemExit(f"Missing master DOCX: {master}")
    translation_count = 0
    editing_count = 0
    for row in active_rows(master):
        assignment = row.cells[2].text.strip()
        if not assignment:
            if not has_english_details(row.cells[1]):
                translation_count += 1
            continue

        roles = [part.strip() for part in assignment.split("/", 1)]
        translator = roles[0]
        editor = roles[1] if len(roles) == 2 else ""
        if editor:
            continue
        key = normalize_name(translator)
        policy = policies.get(key)
        if policy is None:
            label = row.cells[0].text.strip() or "?"
            raise SystemExit(
                f"Unknown translator {translator!r} in {master.name}, row {label}. "
                f"Add them to {CONFIG_PATH.name}."
            )
        if policy.edit_required:
            editing_count += 1
    return translation_count, editing_count


def format_section(title: str, counts: list[tuple[str, int]]) -> str:
    lines = [f"{title}："]
    nonzero = [(programme, count) for programme, count in counts if count]
    if not nonzero:
        lines.append("無")
    else:
        lines.extend(f"{count}集{programme}" for programme, count in nonzero)
    return "\n".join(lines)


def build_report() -> str:
    policies, ready_target = load_config(CONFIG_PATH)
    translations: list[tuple[str, int]] = []
    edits: list[tuple[str, int]] = []
    for programme, master in PROGRAMMES:
        translation_count, editing_count = count_programme(master, policies)
        translations.append((programme, translation_count))
        edits.append((programme, editing_count))
    sections = [
        format_section("待翻譯的節目", translations),
        format_section("待edit的節目", edits),
    ]
    low_queues = [(programme, count) for programme, count in translations if count < ready_target]
    if low_queues:
        warning_lines = ["提醒："]
        warning_lines.extend(
            f"{programme}待翻譯節目不足{ready_target}集，目前{count}集，請再新增{ready_target - count}集。"
            for programme, count in low_queues
        )
        sections.append("\n".join(warning_lines))
    return "\n\n".join(sections)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--no-copy", action="store_true", help="print without calling wl-copy")
    args = parser.parse_args()
    report = build_report()
    print(report)
    if not args.no_copy:
        try:
            subprocess.run(["wl-copy"], input=report, text=True, check=True)
        except FileNotFoundError:
            raise SystemExit("wl-copy is not installed; report was printed but not copied.")
        except subprocess.CalledProcessError as error:
            raise SystemExit(f"wl-copy failed with status {error.returncode}; report was printed but not copied.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
